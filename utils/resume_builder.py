"""
Resume Builder - Dynamic resume tailoring based on job description
Uses python-docx to modify resume template
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from typing import List, Optional
from config import ASSETS_DIR, OUTPUT_DIR


# Placeholder markers in template
PLACEHOLDERS = {
    "skills": "{{SKILLS_SECTION}}",
    "summary": "{{SUMMARY}}",
    "keywords": "{{KEYWORDS}}"
}


def create_template_resume():
    """Create a basic resume template if it doesn't exist"""
    template_path = os.path.join(ASSETS_DIR, "master_resume.docx")
    
    if os.path.exists(template_path):
        return template_path
    
    # Create basic template
    doc = Document()
    
    # Add name (placeholder)
    name_para = doc.add_paragraph()
    name_run = name_para.add_run("YOUR NAME")
    name_run.bold = True
    name_run.font.size = Pt(16)
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Contact info
    contact = doc.add_paragraph("email@example.com | +91-XXXXXXXXXX | LinkedIn: linkedin.com/in/yourprofile")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Summary section
    doc.add_paragraph("PROFESSIONAL SUMMARY", style='Heading 1')
    doc.add_paragraph("{{SUMMARY}}")
    
    # Skills section
    doc.add_paragraph("TECHNICAL SKILLS", style='Heading 1')
    doc.add_paragraph("{{SKILLS_SECTION}}")
    
    # Experience section
    doc.add_paragraph("EXPERIENCE", style='Heading 1')
    doc.add_paragraph("Add your experience here...")
    
    # Education section
    doc.add_paragraph("EDUCATION", style='Heading 1')
    doc.add_paragraph("Add your education here...")
    
    # Projects section
    doc.add_paragraph("PROJECTS", style='Heading 1')
    doc.add_paragraph("Add your projects here...")
    
    # Save template
    doc.save(template_path)
    print(f"✅ Created template resume at: {template_path}")
    
    return template_path


def replace_in_paragraph(paragraph, placeholder: str, replacement: str):
    """Replace placeholder in a paragraph while preserving formatting"""
    if placeholder in paragraph.text:
        # Simple replacement for full paragraph
        inline = paragraph.runs
        for run in inline:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, replacement)


def replace_in_document(doc: Document, placeholder: str, replacement: str):
    """Replace placeholder throughout the document"""
    # Replace in paragraphs
    for para in doc.paragraphs:
        replace_in_paragraph(para, placeholder, replacement)
    
    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_paragraph(para, placeholder, replacement)


def tailor_resume(
    job_description: str,
    company: str,
    role: str,
    keywords: List[str] = None,
    template_path: str = None
) -> Optional[str]:
    """
    Create a tailored resume based on job description.
    
    Args:
        job_description: Full job description text
        company: Company name
        role: Job role/title
        keywords: Pre-extracted keywords (optional)
        template_path: Path to template resume (optional)
    
    Returns:
        Path to generated resume file
    """
    # Use default template if not specified
    if template_path is None:
        template_path = os.path.join(ASSETS_DIR, "master_resume.docx")
    
    # Create template if it doesn't exist
    if not os.path.exists(template_path):
        create_template_resume()
    
    try:
        # Load template
        doc = Document(template_path)
        
        # Generate keywords if not provided
        if keywords is None or len(keywords) == 0:
            # Try to use AI brain for keyword extraction
            try:
                from utils.ai_brain import extract_keywords
                keywords = extract_keywords(job_description)
            except:
                keywords = ["Python", "Problem Solving", "Communication", "Team Work"]
        
        # Format keywords as comma-separated string
        keywords_str = ", ".join(keywords) if keywords else "Various technical skills"
        
        # Generate tailored summary
        summary = f"Motivated professional seeking the {role} position at {company}. " \
                  f"Skilled in {', '.join(keywords[:3]) if len(keywords) >= 3 else 'relevant technologies'}. " \
                  f"Eager to contribute to innovative projects and grow professionally."
        
        # Replace placeholders
        replace_in_document(doc, PLACEHOLDERS["skills"], keywords_str)
        replace_in_document(doc, PLACEHOLDERS["summary"], summary)
        replace_in_document(doc, PLACEHOLDERS["keywords"], keywords_str)
        
        # Create safe filename
        safe_company = "".join(c for c in company if c.isalnum() or c in " -_")[:25]
        safe_role = "".join(c for c in role if c.isalnum() or c in " -_")[:25]
        filename = f"Resume_{safe_company}_{safe_role}.docx".replace(" ", "_")
        
        # Full output path
        output_path = os.path.join(OUTPUT_DIR, filename)
        
        # Ensure output directory exists
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        
        # Save tailored resume
        doc.save(output_path)
        print(f"✅ Tailored resume saved: {output_path}")
        
        return output_path
        
    except Exception as e:
        print(f"❌ Error creating resume: {e}")
        return None


def get_resume_text(docx_path: str) -> str:
    """Extract text from a DOCX file"""
    try:
        doc = Document(docx_path)
        full_text = []
        
        for para in doc.paragraphs:
            full_text.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        
        return "\n".join(full_text)
    except Exception as e:
        print(f"Error reading resume: {e}")
        return ""


def list_generated_resumes() -> List[str]:
    """List all generated resumes"""
    if not os.path.exists(OUTPUT_DIR):
        return []
    
    resumes = []
    for file in os.listdir(OUTPUT_DIR):
        if file.startswith("Resume_") and file.endswith(".docx"):
            resumes.append(os.path.join(OUTPUT_DIR, file))
    
    return resumes


if __name__ == "__main__":
    # Test the resume builder
    print("\n=== Testing Resume Builder ===\n")
    
    # Create template first
    template = create_template_resume()
    print(f"Template: {template}")
    
    # Test tailoring
    test_jd = """
    We are looking for a Python Developer with experience in:
    - Django or FastAPI
    - PostgreSQL
    - AWS or GCP
    - Docker, Kubernetes
    - REST APIs
    
    Nice to have: Machine Learning, Redis, Celery
    """
    
    output = tailor_resume(
        job_description=test_jd,
        company="TechCorp",
        role="Python Developer",
        keywords=["Python", "Django", "FastAPI", "PostgreSQL", "AWS", "Docker"]
    )
    
    if output:
        print(f"\n📄 Resume generated: {output}")
    else:
        print("\n❌ Failed to generate resume")
